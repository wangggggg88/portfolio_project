"""Build outputs/portfolio/attribution_check_<portfolio>.xlsx: Brinson-Fachler + Carino linking in native Excel formulas.

Every effect is a live formula on month-end prices, so the workbook is an
independent re-implementation of src/attribution.py. The Check sheet compares
the formula results with the Python values written by `python -m src.run_v2`.

    python scripts/build_attribution_excel.py            # build (needs openpyxl)
    python scripts/build_attribution_excel.py --verify   # also recalculate in LibreOffice and compare
"""
import argparse
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from src.quality import validate  # noqa: E402

HEAD = Font(bold=True)
FILL = PatternFill("solid", fgColor="F0EFEA")
INPUT = Font(color="1F4E9E")


def build(config="config/full.json", portfolio="Main"):
    cfg = json.loads((ROOT / config).read_text(encoding="utf-8"))
    frame = pd.read_csv(ROOT / cfg["prices"])
    cal = pd.read_csv(ROOT / cfg["calendar"]).date
    prices, _ = validate(frame, cal, list(cfg["securities"]), cfg["anchor"], cfg["end"])
    me = prices.groupby(prices.index.to_period("M")).tail(1)
    tick = list(prices.columns)
    n = len(tick)
    seg = {t: ("Equity" if "equity" in v["asset_class"].lower() else "Fixed income") for t, v in cfg["securities"].items()}
    wb = Workbook()

    # ---- Prices (inputs) ----
    ws = wb.active; ws.title = "Prices"
    ws.append(["Month-end adjusted close (input, from data/snapshot/prices.csv)"]); ws["A1"].font = HEAD
    ws.append(["date"] + tick)
    for c in ws[2]: c.font = HEAD; c.fill = FILL
    for d, row in me.iterrows():
        ws.append([d.date()] + [float(x) for x in row])
        for c in ws[ws.max_row][1:]: c.font = INPUT
    first_p, last_p = 3, ws.max_row          # row 3 = anchor 2020-12-31
    ws.column_dimensions["A"].width = 12

    # ---- Weights (inputs), one column per ticker so formulas need no array entry ----
    ww = wb.create_sheet("Weights")
    ww.append(["ticker"] + tick)
    ww.append(["segment"] + [seg[t] for t in tick])
    ww.append(["portfolio_w"] + [cfg["portfolios"][portfolio].get(t, 0.0) for t in tick])
    ww.append(["benchmark_w"] + [cfg["portfolios"]["Benchmark"].get(t, 0.0) for t in tick])
    ww.append(["is_equity"] + [f'=IF({L(2 + j)}2="Equity",1,0)' for j in range(n)])
    for c in ww[1]: c.font = HEAD; c.fill = FILL
    for rr in (3, 4):
        for c in ww[rr][1:]: c.font = INPUT
    W = f"Weights!$B$3:${L(1 + n)}$3"; B = f"Weights!$B$4:${L(1 + n)}$4"; E = f"Weights!$B$5:${L(1 + n)}$5"
    ww.append([]); ww.append(["Portfolio", portfolio]); ww.append(["Benchmark", "Benchmark"])
    ww.append(["w_p equity", f"=SUMPRODUCT({W},{E})"]); ww.append(["w_p fixed income", f"=SUM({W})-B9"])
    ww.append(["w_b equity", f"=SUMPRODUCT({B},{E})"]); ww.append(["w_b fixed income", f"=SUM({B})-B11"])
    WPE, WPF, WBE, WBF = "Weights!$B$9", "Weights!$B$10", "Weights!$B$11", "Weights!$B$12"
    ww.column_dimensions["A"].width = 18

    # ---- Brinson (formulas) ----
    bs = wb.create_sheet("Brinson")
    heads = ["month"] + [f"R_{t}" for t in tick] + [
        "Rp", "Rb", "rp_Eq", "rp_FI", "rb_Eq", "rb_FI",
        "alloc_Eq", "alloc_FI", "sel_Eq", "sel_FI", "int_Eq", "int_FI", "total_effects", "check_vs_active",
        "k_t", "cum_p", "cum_b"]
    bs.append(heads)
    for c in bs[1]: c.font = HEAD; c.fill = FILL; c.alignment = Alignment(horizontal="center")
    col = {h: i + 1 for i, h in enumerate(heads)}
    C = lambda h: L(col[h])
    rcols = f"{L(2)}{{r}}:{L(1 + n)}{{r}}"
    for i in range(first_p + 1, last_p + 1):
        r = bs.max_row + 1
        bs.cell(r, 1).value = f"=TEXT(Prices!A{i},\"yyyy-mm\")"
        for j in range(n):
            p = L(2 + j)
            bs.cell(r, 2 + j).value = f"=Prices!{p}{i}/Prices!{p}{i - 1}-1"
        rng = rcols.format(r=r)
        bs.cell(r, col["Rp"]).value = f"=SUMPRODUCT({rng},{W})"
        bs.cell(r, col["Rb"]).value = f"=SUMPRODUCT({rng},{B})"
        bs.cell(r, col["rp_Eq"]).value = f"=SUMPRODUCT({rng},{W},{E})/{WPE}"
        bs.cell(r, col["rp_FI"]).value = f"=SUMPRODUCT({rng},{W},1-{E})/{WPF}"
        bs.cell(r, col["rb_Eq"]).value = f"=SUMPRODUCT({rng},{B},{E})/{WBE}"
        bs.cell(r, col["rb_FI"]).value = f"=SUMPRODUCT({rng},{B},1-{E})/{WBF}"
        Rb = f"{C('Rb')}{r}"
        bs.cell(r, col["alloc_Eq"]).value = f"=({WPE}-{WBE})*({C('rb_Eq')}{r}-{Rb})"
        bs.cell(r, col["alloc_FI"]).value = f"=({WPF}-{WBF})*({C('rb_FI')}{r}-{Rb})"
        bs.cell(r, col["sel_Eq"]).value = f"={WBE}*({C('rp_Eq')}{r}-{C('rb_Eq')}{r})"
        bs.cell(r, col["sel_FI"]).value = f"={WBF}*({C('rp_FI')}{r}-{C('rb_FI')}{r})"
        bs.cell(r, col["int_Eq"]).value = f"=({WPE}-{WBE})*({C('rp_Eq')}{r}-{C('rb_Eq')}{r})"
        bs.cell(r, col["int_FI"]).value = f"=({WPF}-{WBF})*({C('rp_FI')}{r}-{C('rb_FI')}{r})"
        bs.cell(r, col["total_effects"]).value = f"=SUM({C('alloc_Eq')}{r}:{C('int_FI')}{r})"
        bs.cell(r, col["check_vs_active"]).value = f"=ABS({C('total_effects')}{r}-({C('Rp')}{r}-{Rb}))"
        Rp = f"{C('Rp')}{r}"
        bs.cell(r, col["k_t"]).value = f"=IF(ABS({Rp}-{Rb})<1E-15,1/(1+{Rp}),(LN(1+{Rp})-LN(1+{Rb}))/({Rp}-{Rb}))"
        prev = "1" if r == 2 else f"{C('cum_p')}{r - 1}"
        prevb = "1" if r == 2 else f"{C('cum_b')}{r - 1}"
        bs.cell(r, col["cum_p"]).value = f"={prev}*(1+{Rp})"
        bs.cell(r, col["cum_b"]).value = f"={prevb}*(1+{Rb})"
    last = bs.max_row
    bs.freeze_panes = "B2"
    for k in range(1, len(heads) + 1):
        bs.column_dimensions[L(k)].width = 11
        for rr in range(2, last + 1):
            bs.cell(rr, k).number_format = "0.0000%" if k > 1 else "@"

    # ---- Summary & Check ----
    sm = wb.create_sheet("Check", 0)
    sm.append([f"Brinson-Fachler attribution, {portfolio} vs Benchmark, Carino-linked (live formulas)"]); sm["A1"].font = Font(bold=True, size=12)
    sm.append(["Months", f"=COUNT(Brinson!{C('Rp')}2:{C('Rp')}{last})"])
    sm.append(["Cumulative portfolio return", f"=Brinson!{C('cum_p')}{last}-1"])
    sm.append(["Cumulative benchmark return", f"=Brinson!{C('cum_b')}{last}-1"])
    sm.append(["Active return", "=B3-B4"])
    sm.append(["K (Carino total factor)", "=(LN(1+B3)-LN(1+B4))/(B3-B4)"])
    sm.append([])
    sm.append(["Linked effect", "Excel formula", "Python (reports/v2)", "Abs difference", "Result"])
    for c in sm[8]: c.font = HEAD; c.fill = FILL
    kcol = f"Brinson!${C('k_t')}$2:${C('k_t')}${last}"
    py = pd.read_csv(ROOT / "reports/v2" / f"attribution_{portfolio}_by_segment.csv", index_col=0)
    items = [("Allocation, equity", "alloc_Eq", py.loc["Equity", "allocation"]),
             ("Allocation, fixed income", "alloc_FI", py.loc["Fixed income", "allocation"]),
             ("Selection, equity", "sel_Eq", py.loc["Equity", "selection"]),
             ("Selection, fixed income", "sel_FI", py.loc["Fixed income", "selection"]),
             ("Interaction, equity", "int_Eq", py.loc["Equity", "interaction"]),
             ("Interaction, fixed income", "int_FI", py.loc["Fixed income", "interaction"])]
    for label, h, value in items:
        r = sm.max_row + 1
        sm.append([label, f"=SUMPRODUCT(Brinson!{C(h)}2:{C(h)}{last},{kcol})/$B$6", float(value), f"=ABS(B{r}-C{r})",
                   f'=IF(D{r}<=1E-10,"PASS","FAIL")'])
    first_item, last_item = 9, sm.max_row
    r = sm.max_row + 1
    sm.append(["Sum of linked effects", f"=SUM(B{first_item}:B{last_item})", f"=SUM(C{first_item}:C{last_item})",
               f"=ABS(B{r}-B5)", f'=IF(D{r}<=1E-10,"PASS","FAIL")'])
    sm.append(["Max monthly |effects - active|", f"=MAX(Brinson!{C('check_vs_active')}2:{C('check_vs_active')}{last})", "", "",
               f'=IF(B{r + 1}<=1E-12,"PASS","FAIL")'])
    sm.append([]); sm.append(["Overall", f'=IF(COUNTIF(E{first_item}:E{r + 1},"PASS")={r + 2 - first_item},"ALL PASS","CHECK FAILURES")'])
    sm["A" + str(sm.max_row)].font = HEAD; sm["B" + str(sm.max_row)].font = HEAD
    for rr in range(3, sm.max_row + 1):
        for cc in (2, 3):
            sm.cell(rr, cc).number_format = "0.0000%"
        sm.cell(rr, 4).number_format = "0.00E+00"
    sm["B2"].number_format = "0"; sm["B6"].number_format = "0.000000"
    sm.column_dimensions["A"].width = 32
    for c in "BCDE": sm.column_dimensions[c].width = 18
    sm.append([]); sm.append(["Blue numbers on Prices/Weights are inputs; everything else is a formula. Retrospective simulation, not investment advice."])
    out = ROOT / f"outputs/portfolio/attribution_check_{portfolio}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def verify(path, portfolio):
    """Recalculate with LibreOffice headless and read back cached values."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise SystemExit("LibreOffice not found; open the workbook in Excel and read the Check sheet instead")
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp, str(path)],
                       check=True, capture_output=True)
        recalculated = Path(tmp) / path.name
        wbv = load_workbook(recalculated, data_only=True)
        ws = wbv["Check"]
        rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2, max_row=ws.max_row)]
        shutil.copy(recalculated, path)  # keep formulas plus cached values for previewers
    result = {"workbook": path.name, "portfolio": portfolio,
              "engine": "LibreOffice Calc (headless recalculation)", "desktop_excel_executed": False,
              "rows": [r for r in rows if r and r[0]]}
    overall = [r for r in rows if r and r[0] == "Overall"][0][1]
    result["overall"] = overall
    print(path.name, overall)
    if overall != "ALL PASS":
        raise SystemExit("Workbook checks failed: " + json.dumps(result, default=str))
    return result


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--verify", action="store_true")
    a = ap.parse_args()
    results = []
    for name in ["Main", "Conservative"]:
        p = build(portfolio=name)
        print("built", p)
        if a.verify:
            results.append(verify(p, name))
    if a.verify:
        (ROOT / "reports/v2/excel_attribution_verification.json").write_text(
            json.dumps(results, indent=2, default=str), encoding="utf-8")
